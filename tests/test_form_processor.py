"""
Test suite for FormProcessor class.

Tests for:
- Random response generation
- File data loading (CSV, JSON)
- User edit application
"""
import pytest
import json
import tempfile
import os
import random
from io import BytesIO
from app.core.form_processor import BRANCH_SUBMIT_SENTINEL, FormProcessor
from app.models import Form, ResponseConfig, Page, Question, AnswerConfig, AnswerOption
from datetime import datetime
import app.main_routes as main_routes
from openpyxl import Workbook


@pytest.fixture
def sample_form():
    """Create a sample form with various question types."""
    return Form(
        id="test_form_001",
        title="Test Form",
        description="Test description",
        url="https://example.com/form",
        created_at=datetime.now(),
        last_used=None,
        response_config=ResponseConfig(pages=[
            Page(
                page_id="page_1",
                questions=[
                    Question(
                        question_id="q1",
                        type="input_text",
                        text="Your name?",
                        answer_config=AnswerConfig(fill_percentage=100, answers=[], options=None)
                    ),
                    Question(
                        question_id="q2",
                        type="multiple_choice",
                        text="Favorite color?",
                        answer_config=AnswerConfig(
                            fill_percentage=100,
                            answers=None,
                            options=[
                                AnswerOption(text="Red", percentage=50.0),
                                AnswerOption(text="Blue", percentage=30.0),
                                AnswerOption(text="Green", percentage=20.0),
                            ]
                        )
                    ),
                    Question(
                        question_id="q3",
                        type="checkbox",
                        text="Select hobbies",
                        answer_config=AnswerConfig(
                            fill_percentage=100,
                            answers=None,
                            options=[
                                AnswerOption(text="Reading", percentage=50.0),
                                AnswerOption(text="Sports", percentage=50.0),
                            ]
                        )
                    ),
                ]
            )
        ]),
        submissions=None
    )


@pytest.fixture
def processor(sample_form):
    """Create a FormProcessor instance."""
    return FormProcessor(sample_form)


class TestRandomResponses:
    """Tests for random response generation."""

    def test_generate_random_responses_returns_dict(self, processor):
        responses = processor.generate_random_responses()
        assert isinstance(responses, dict)

    def test_generate_random_responses_has_all_questions(self, processor):
        responses = processor.generate_random_responses(fill_percentage=100)
        assert "q1" in responses
        assert "q2" in responses
        assert "q3" in responses

    def test_text_response_is_string(self, processor):
        responses = processor.generate_random_responses()
        assert isinstance(responses.get("q1"), str)

    def test_multiple_choice_response_is_string(self, processor):
        responses = processor.generate_random_responses()
        q2_response = responses.get("q2")
        assert q2_response in ["Red", "Blue", "Green"]

    def test_checkbox_response_is_list(self, processor):
        responses = processor.generate_random_responses()
        q3_response = responses.get("q3")
        assert isinstance(q3_response, list)
        for item in q3_response:
            assert item in ["Reading", "Sports"]

    def test_question_fill_percentage_zero_skips_question(self, processor):
        q1 = processor.form.response_config.pages[0].questions[0]
        q1.answer_config.fill_percentage = 0

        responses = processor.generate_random_responses(fill_percentage=100)

        assert "q1" not in responses

    def test_text_question_uses_configured_answers(self, processor):
        q1 = processor.form.response_config.pages[0].questions[0]
        q1.answer_config.answers = ["Alice", "Bob"]

        responses = processor.generate_random_responses(fill_percentage=100)

        assert responses["q1"] in ["Alice", "Bob"]

    def test_multiple_choice_uses_weighted_percentages(self, processor):
        q2 = processor.form.response_config.pages[0].questions[1]
        q2.answer_config.options = [
            AnswerOption(text="Red", percentage=100),
            AnswerOption(text="Blue", percentage=0),
        ]

        responses = processor.generate_random_responses(fill_percentage=100)

        assert responses["q2"] == "Red"

    def test_multiple_choice_zero_weights_falls_back_to_random_option(self, processor, monkeypatch):
        q2 = processor.form.response_config.pages[0].questions[1]
        q2.answer_config.options = [
            AnswerOption(text="Red", percentage=0),
            AnswerOption(text="Blue", percentage=0),
        ]
        monkeypatch.setattr(random, "choice", lambda options: options[1])

        responses = processor.generate_random_responses(fill_percentage=100)

        assert responses["q2"] == "Blue"

    def test_checkbox_uses_independent_option_percentages(self, processor):
        q3 = processor.form.response_config.pages[0].questions[2]
        q3.answer_config.options = [
            AnswerOption(text="Reading", percentage=100),
            AnswerOption(text="Sports", percentage=0),
        ]

        responses = processor.generate_random_responses(fill_percentage=100)

        assert responses["q3"] == ["Reading"]

    def test_checkbox_zero_weights_keeps_existing_random_fallback(self, processor, monkeypatch):
        q3 = processor.form.response_config.pages[0].questions[2]
        q3.answer_config.options = [
            AnswerOption(text="Reading", percentage=0),
            AnswerOption(text="Sports", percentage=0),
        ]
        monkeypatch.setattr(random, "randint", lambda start, end: 1)
        monkeypatch.setattr(random, "sample", lambda options, count: [options[0]])

        responses = processor.generate_random_responses(fill_percentage=100)

        assert responses["q3"] == ["Reading"]

    def test_prefill_responses_use_configured_text_rows(self, processor):
        q1 = processor.form.response_config.pages[0].questions[0]
        q1.answer_config.answers = ["Alice", "Bob"]

        responses = processor.generate_prefill_responses(3)

        assert [row["q1"] for row in responses] == ["Alice", "Bob", "Alice"]

    def test_prefill_responses_materialize_single_option_distribution(self, processor):
        q2 = processor.form.response_config.pages[0].questions[1]
        q2.answer_config.options = [
            AnswerOption(text="Red", percentage=75),
            AnswerOption(text="Blue", percentage=25),
        ]

        responses = processor.generate_prefill_responses(4)

        assert [row["q2"] for row in responses] == ["Red", "Red", "Red", "Blue"]

    def test_prefill_responses_materialize_rank_as_single_value(self, sample_form):
        sample_form.response_config.pages[0].questions = [
            Question(
                question_id="q_rank",
                entry_id="entry.333",
                type="rank",
                text="Rate this",
                answer_config=AnswerConfig(
                    fill_percentage=None,
                    answers=None,
                    options=[
                        AnswerOption(text="1", percentage=25),
                        AnswerOption(text="2", percentage=75),
                    ],
                ),
            )
        ]
        processor = FormProcessor(sample_form)

        responses = processor.generate_prefill_responses(4)

        assert [row["q_rank"] for row in responses] == ["1", "2", "2", "2"]

    def test_prefill_responses_materialize_checkbox_repeated_values(self, processor):
        q3 = processor.form.response_config.pages[0].questions[2]
        q3.answer_config.options = [
            AnswerOption(text="Reading", percentage=50),
            AnswerOption(text="Sports", percentage=50),
        ]

        responses = processor.generate_prefill_responses(4)

        assert [row["q3"] for row in responses] == [
            ["Reading", "Sports"],
            ["Reading", "Sports"],
            ["Reading"],
            ["Reading"],
        ]

    def test_prefill_responses_skip_later_pages_for_submit_branch(self, sample_form):
        sample_form.response_config.pages = [
            Page(
                page_id="page_1",
                questions=[
                    Question(
                        question_id="q_branch",
                        type="multiple_choice",
                        text="Continue?",
                        answer_config=AnswerConfig(
                            fill_percentage=100,
                            answers=None,
                            options=[
                                AnswerOption(text="Continue", percentage=50),
                                AnswerOption(
                                    text="Submit now",
                                    percentage=50,
                                    next_page_id=BRANCH_SUBMIT_SENTINEL,
                                ),
                            ],
                        ),
                    )
                ],
            ),
            Page(
                page_id="page_2",
                questions=[
                    Question(
                        question_id="q_followup",
                        type="input_text",
                        text="Follow-up answer",
                        answer_config=AnswerConfig(
                            fill_percentage=100,
                            answers=["Follow-up value"],
                            options=None,
                        ),
                    )
                ],
            ),
        ]
        processor = FormProcessor(sample_form)

        responses = processor.generate_prefill_responses(2)

        assert responses[0]["q_branch"] == "Continue"
        assert responses[0]["q_followup"] == "Follow-up value"
        assert responses[1]["q_branch"] == "Submit now"
        assert "q_followup" not in responses[1]


class TestLoadDataFromFile:
    """Tests for file loading functionality."""

    def test_load_csv_file(self, processor):
        csv_content = "q1,q2,q3\nJohn,Red,Reading\nJane,Blue,\"Reading,Sports\""
        with tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False, encoding='utf-8') as f:
            f.write(csv_content)
            csv_path = f.name

        try:
            responses_list = processor.load_data_from_file(csv_path)
            assert len(responses_list) == 2
            assert responses_list[0]["q1"] == "John"
            assert responses_list[0]["q2"] == "Red"
            assert responses_list[1]["q1"] == "Jane"
        finally:
            os.unlink(csv_path)

    def test_load_json_file(self, processor):
        json_content = [
            {"q1": "Alice", "q2": "Green", "q3": ["Sports"]},
            {"q1": "Bob", "q2": "Red", "q3": ["Reading", "Sports"]}
        ]
        with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False, encoding='utf-8') as f:
            json.dump(json_content, f)
            json_path = f.name

        try:
            responses_list = processor.load_data_from_file(json_path)
            assert len(responses_list) == 2
            assert responses_list[0]["q1"] == "Alice"
            assert responses_list[1]["q3"] == ["Reading", "Sports"]
        finally:
            os.unlink(json_path)

    def test_load_with_mapping(self, processor):
        csv_content = "name,color,hobbies\nJohn,Red,Reading"
        with tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False, encoding='utf-8') as f:
            f.write(csv_content)
            csv_path = f.name

        try:
            mapping = {"name": "q1", "color": "q2", "hobbies": "q3"}
            responses_list = processor.load_data_from_file(csv_path, mapping)
            assert len(responses_list) == 1
            assert responses_list[0]["q1"] == "John"
            assert responses_list[0]["q2"] == "Red"
        finally:
            os.unlink(csv_path)

    def test_load_xlsx_file(self, processor):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["q1", "q2", "q3"])
        worksheet.append(["Alice", "Green", True])
        worksheet.append(["Bob", "Red", 7])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            xlsx_path = f.name

        try:
            workbook.save(xlsx_path)
            responses_list = processor.load_data_from_file(xlsx_path)
            assert len(responses_list) == 2
            assert responses_list[0] == {"q1": "Alice", "q2": "Green", "q3": True}
            assert responses_list[1] == {"q1": "Bob", "q2": "Red", "q3": 7}
        finally:
            workbook.close()
            os.unlink(xlsx_path)

    def test_load_xlsx_with_mapping(self, processor):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["name", "color", "hobbies"])
        worksheet.append(["Dana", "Blue", "Reading"])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            xlsx_path = f.name

        try:
            workbook.save(xlsx_path)
            mapping = {"name": "q1", "color": "q2", "hobbies": "q3"}
            responses_list = processor.load_data_from_file(xlsx_path, mapping)
            assert responses_list == [{"q1": "Dana", "q2": "Blue", "q3": "Reading"}]
        finally:
            workbook.close()
            os.unlink(xlsx_path)

    def test_load_xlsx_skips_empty_rows_and_cells(self, processor):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["q1", "q2", "q3"])
        worksheet.append(["Alice", None, True])
        worksheet.append([None, None, None])
        worksheet.append(["Bob", "", 12])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            xlsx_path = f.name

        try:
            workbook.save(xlsx_path)
            responses_list = processor.load_data_from_file(xlsx_path)
            assert responses_list == [
                {"q1": "Alice", "q3": True},
                {"q1": "Bob", "q3": 12},
            ]
        finally:
            workbook.close()
            os.unlink(xlsx_path)

    def test_load_nonexistent_file_raises(self, processor):
        with pytest.raises(FileNotFoundError):
            processor.load_data_from_file("/nonexistent/path/file.csv")

    def test_load_unsupported_format_raises(self, processor):
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as f:
            file_path = f.name

        try:
            with pytest.raises(ValueError, match="Use .csv, .json, or .xlsx"):
                processor.load_data_from_file(file_path)
        finally:
            os.unlink(file_path)

    def test_load_data_route_accepts_xlsx(self, client, sample_form, monkeypatch):
        class Storage:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc_value, traceback):
                return None

            def _load_form(self, form_id):
                return sample_form

        def load_data_from_file(self, file_path, mapping=None):
            assert file_path.endswith(".xlsx")
            return [{"q1": "Alice"}]

        monkeypatch.setattr(main_routes, "get_storage_service", lambda: Storage())
        monkeypatch.setattr(FormProcessor, "load_data_from_file", load_data_from_file)

        response = client.post(
            "/load_data",
            data={
                "form_id": sample_form.id,
                "answer_file": (BytesIO(b"not used by mocked parser"), "answers.xlsx"),
            },
            content_type="multipart/form-data",
        )

        assert response.status_code == 200
        assert response.get_json()["rows_loaded"] == 1


class TestApplyUserEdits:
    """Tests for applying user edits to form configuration."""

    def test_apply_fill_percentage_edit(self, processor):
        edits = {
            "q1": {"fill_percentage": 50}
        }
        processor.apply_user_edits(edits)

        q1 = processor.form.response_config.pages[0].questions[0]
        assert q1.answer_config.fill_percentage == 50

    def test_apply_answers_edit(self, processor):
        edits = {
            "q1": {"answers": ["Custom Answer 1", "Custom Answer 2"]}
        }
        processor.apply_user_edits(edits)

        q1 = processor.form.response_config.pages[0].questions[0]
        assert q1.answer_config.answers == ["Custom Answer 1", "Custom Answer 2"]

    def test_apply_options_percentage_edit(self, processor):
        edits = {
            "q2": {
                "options": [
                    {"text": "Red", "percentage": 80},
                    {"text": "Blue", "percentage": 10},
                    {"text": "Green", "percentage": 10},
                ]
            }
        }
        processor.apply_user_edits(edits)

        q2 = processor.form.response_config.pages[0].questions[1]
        options = {opt.text: opt.percentage for opt in q2.answer_config.options}
        assert options["Red"] == 80
        assert options["Blue"] == 10

    def test_apply_options_next_page_edit(self, processor):
        edits = {
            "q2": {
                "options": [
                    {"text": "Blue", "percentage": 30, "next_page_id": BRANCH_SUBMIT_SENTINEL},
                ]
            }
        }

        processor.apply_user_edits(edits)

        q2 = processor.form.response_config.pages[0].questions[1]
        options = {opt.text: opt.next_page_id for opt in q2.answer_config.options}
        assert options["Blue"] == BRANCH_SUBMIT_SENTINEL

    def test_apply_edit_to_nonexistent_question_is_ignored(self, processor):
        edits = {
            "nonexistent_q": {"fill_percentage": 50}
        }
        processor.apply_user_edits(edits)
